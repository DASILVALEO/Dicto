"""
Generador de comprobantes de retencion con openpyxl (mas confiable que gridcraft-cli).

Usa el template real del usuario pero crea un archivo NUEVO basado en la
hoja activa del template, con los datos del comprobante anterior eliminados
para que la nueva factura tenga campos limpios.
"""
import os
import sys
import json
import shutil
from pathlib import Path
from datetime import datetime
from typing import Optional
import openpyxl

TEMPLATES_DIR = Path(__file__).parent / "templates"

# Celdas conocidas en el template de comprobante IVA
# (basado en el analisis de las celdas reales)
CELDS_IVA = {
    "numero_comprobante": "M5",
    "fecha_emision": "Q5",
    "fecha_entrega": "Q9",
    "sujeto_nombre": "A18",
    "sujeto_rif": "I18",
    "sujeto_direccion": "A21",
    # Fila de items (la primera, fila 27). Las demas se limpian.
    "item_operacion": "A27",
    "item_fecha": "B27",
    "item_factura": "C27",
    "item_control": "D27",
    "item_tipo": "H27",
    "item_total": "I27",
    "item_base": "L27",
    "item_alicuota": "M27",
    "item_impuesto": "N27",
    "item_retencion_pct": "O27",
    "item_retenido": "P27",
    "item_neto": "Q27",
}


def find_active_sheet(wb):
    """Encuentra la hoja activa (ultima hoja visible con data)."""
    candidates = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.sheet_state == "visible" and ws.max_row > 5:
            con_data = sum(1 for row in ws.iter_rows() for c in row if c.value)
            if con_data > 50:
                candidates.append((sn, con_data, ws))
    # Tomar la ultima (mas reciente)
    if candidates:
        return candidates[-1][2]
    return None


def crear_comprobante_iva(datos: dict, output_dir: Path) -> Optional[Path]:
    """Crea un comprobante de retencion IVA."""
    template = TEMPLATES_DIR / "COMPROBANTE DE RETENCION IVA -AUTOMERCADO-SEPT-OCTUBRE  2026.xlsx"
    if not template.exists():
        return None
    # Nombre de archivo de salida
    rif_clean = datos['rif_proveedor'].replace('-', '').replace('.', '')
    factura_clean = datos['numero_factura'].replace('/', '-').replace('\\', '-')
    output = output_dir / f"Comprobante_IVA_{rif_clean}_{factura_clean}.xlsx"
    # Copiar el template
    shutil.copy2(template, output)
    # Abrir con openpyxl
    wb = openpyxl.load_workbook(output)
    ws = find_active_sheet(wb)
    if not ws:
        return None
    # Limpiar filas de items anteriores (fila 27 en adelante)
    # Ojo: las celdas pueden estar fusionadas (merged), no se pueden escribir
    merged_ranges = list(ws.merged_cells.ranges)
    for row in range(27, 50):
        for col in range(1, 20):
            cell = ws.cell(row=row, column=col)
            # Saltar si la celda es parte de un merge (no la principal)
            is_merged = False
            for mr in merged_ranges:
                if cell.coordinate in mr and cell.coordinate != mr.coord.split(":")[0]:
                    is_merged = True
                    break
            if not is_merged:
                try:
                    cell.value = None
                except AttributeError:
                    pass
    # Sobrescribir datos del comprobante
    ws[CELDS_IVA["numero_comprobante"]] = datos.get("numero_comprobante") or \
        f"RET-{datetime.now().strftime('%Y%m%d%H%M%S')}"
    # Fechas: openpyxl necesita objetos datetime
    for cell_key, data_key in [
        ("fecha_emision", "fecha_emision"),
        ("fecha_entrega", "fecha_recepcion"),
    ]:
        fecha_str = datos.get(data_key, "")
        if fecha_str:
            try:
                fecha = datetime.strptime(fecha_str, "%d/%m/%Y")
                ws[CELDS_IVA[cell_key]] = fecha
                ws[CELDS_IVA[cell_key]].number_format = "DD/MM/YYYY"
            except ValueError:
                ws[CELDS_IVA[cell_key]] = fecha_str
    # Datos del proveedor (sujeto retenido)
    ws[CELDS_IVA["sujeto_nombre"]] = datos.get("nombre_proveedor", "")
    ws[CELDS_IVA["sujeto_rif"]] = datos.get("rif_proveedor", "")
    if datos.get("direccion_proveedor"):
        ws[CELDS_IVA["sujeto_direccion"]] = datos["direccion_proveedor"]
    # Item (la unica fila, operacion 001)
    ws[CELDS_IVA["item_operacion"]] = "001"
    ws[CELDS_IVA["item_fecha"]] = _parse_date(datos.get("fecha_emision"))
    ws[CELDS_IVA["item_factura"]] = datos.get("numero_factura", "")
    ws[CELDS_IVA["item_control"]] = datos.get("numero_control", "")
    ws[CELDS_IVA["item_tipo"]] = "1 - REG"
    # Totales (numericos)
    def to_float(v):
        if not v:
            return 0
        try:
            return float(str(v).replace(",", "").replace(" ", ""))
        except ValueError:
            return 0
    ws[CELDS_IVA["item_total"]] = to_float(datos.get("total"))
    ws[CELDS_IVA["item_base"]] = to_float(datos.get("subtotal"))
    ws[CELDS_IVA["item_alicuota"]] = 0.16  # 16% estandar
    ws[CELDS_IVA["item_impuesto"]] = to_float(datos.get("iva"))
    retencion_pct = (datos.get("porcentaje_iva_retenido") or 75) / 100
    ws[CELDS_IVA["item_retencion_pct"]] = retencion_pct
    ws[CELDS_IVA["item_retenido"]] = to_float(datos.get("monto_retencion"))
    neto = to_float(datos.get("total_neto"))
    if not neto:
        neto = to_float(datos.get("total")) - to_float(datos.get("monto_retencion"))
    ws[CELDS_IVA["item_neto"]] = neto
    # Formato de numeros
    for cell_key in ["item_total", "item_base", "item_impuesto", "item_retenido", "item_neto"]:
        ws[CELDS_IVA[cell_key]].number_format = "#,##0.00"
    ws[CELDS_IVA["item_alicuota"]].number_format = "0.00%"
    ws[CELDS_IVA["item_retencion_pct"]].number_format = "0.00%"
    # Guardar
    wb.save(output)
    return output


def crear_comprobante_islr(datos: dict, output_dir: Path) -> Optional[Path]:
    """Crea un comprobante de retencion ISLR."""
    template = TEMPLATES_DIR / "COMPROBANTE DE RETENCIONES DE ISLR-AUTOMERCADO 2026.xlsx"
    if not template.exists():
        return None
    rif_clean = datos['rif_proveedor'].replace('-', '').replace('.', '')
    factura_clean = datos['numero_factura'].replace('/', '-').replace('\\', '-')
    output = output_dir / f"Comprobante_ISLR_{rif_clean}_{factura_clean}.xlsx"
    shutil.copy2(template, output)
    # Por ahora solo copiamos el template (es muy complejo personalizarlo)
    return output


def agregar_a_libro_compras(datos: dict, libro_path: Path) -> bool:
    """Agrega una fila al libro de compras usando openpyxl."""
    if not libro_path.exists():
        return False
    try:
        wb = openpyxl.load_workbook(libro_path)
    except Exception as e:
        print(f"Error abriendo libro: {e}")
        return False
    # Buscar la hoja principal de compras
    hoja_compras = None
    for sn in wb.sheetnames:
        if "COMPRAS" in sn.upper() and "Hidden" not in str(wb[sn].sheet_state):
            if wb[sn].sheet_state == "visible" and wb[sn].max_row > 10:
                hoja_compras = wb[sn]
                break
    if not hoja_compras:
        # Tomar la primera hoja visible
        for sn in wb.sheetnames:
            if wb[sn].sheet_state == "visible" and wb[sn].max_row > 10:
                hoja_compras = wb[sn]
                break
    if not hoja_compras:
        return False
    # Encontrar la primera fila vacia (a partir de fila 11)
    fila = 11
    while fila <= hoja_compras.max_row + 1:
        if hoja_compras.cell(row=fila, column=1).value is None:
            break
        fila += 1
    # Calcular siguiente numero de operacion
    next_op = 1
    for f in range(2, fila):
        v = hoja_compras.cell(row=f, column=1).value
        if v and str(v).isdigit():
            next_op = int(v) + 1
    # Tipo de transaccion
    tipo = "1 - REG"
    if datos.get("tipo_servicio") == "servicios":
        tipo = "2 - SERV"
    elif datos.get("tipo_servicio") == "honorarios":
        tipo = "3 - HON"
    # Escribir
    def to_float(v):
        if not v:
            return 0
        try:
            return float(str(v).replace(",", ""))
        except ValueError:
            return 0
    hoja_compras.cell(row=fila, column=1, value=f"{next_op:03d}")
    hoja_compras.cell(row=fila, column=2, value=_parse_date(datos.get("fecha_emision")))
    hoja_compras.cell(row=fila, column=3, value=datos.get("rif_proveedor", ""))
    hoja_compras.cell(row=fila, column=4, value=datos.get("nombre_proveedor", ""))
    hoja_compras.cell(row=fila, column=5, value=f"COM-{datetime.now().strftime('%Y%m%d%H%M%S')}")
    hoja_compras.cell(row=fila, column=6, value=datos.get("numero_factura", ""))
    hoja_compras.cell(row=fila, column=7, value=datos.get("numero_control", ""))
    hoja_compras.cell(row=fila, column=8, value=to_float(datos.get("total")))
    hoja_compras.cell(row=fila, column=11, value=to_float(datos.get("subtotal")))
    hoja_compras.cell(row=fila, column=12, value=tipo)
    # Formato
    for col in [8, 11]:
        hoja_compras.cell(row=fila, column=col).number_format = "#,##0.00"
    # Guardar
    wb.save(libro_path)
    return True


def _parse_date(date_str):
    """Parsea fecha DD/MM/YYYY a datetime."""
    if not date_str:
        return None
    try:
        return datetime.strptime(date_str, "%d/%m/%Y")
    except ValueError:
        return date_str
