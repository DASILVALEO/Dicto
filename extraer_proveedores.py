"""Extrae el catalogo de proveedores (version optimizada)."""
import openpyxl
import json
from pathlib import Path

FILE = Path(r"C:\Users\USUARIO\Desktop\app\EXCEL-RETENCIONES\EXCEL-RETENCIONES\PROVEEDORES 2021+.xlsx")
OUT = Path(r"C:\Users\USUARIO\Desktop\app\consultas\voz-a-texto\proveedores_catalogo.json")


def main():
    print(f"Abriendo {FILE.name}...")
    wb = openpyxl.load_workbook(FILE, data_only=True)
    # Solo la hoja PROVEEDORES
    ws = wb["PROVEEDORES"]
    print(f"Hoja: PROVEEDORES ({ws.max_row} filas)")

    proveedores = []
    for row in range(3, ws.max_row + 1):  # F1 es header de titulo, F2 es header de columnas
        rif = ws.cell(row=row, column=2).value
        nombre = ws.cell(row=row, column=3).value
        if not rif or not nombre:
            continue
        rif_str = str(rif).strip().upper()
        # Limpiar duplicados en el RIF
        rif_clean = rif_str.split(" ")[0] if " " in rif_str else rif_str
        # Validar formato basico
        if not (rif_clean[0:2] in ("J-", "V-", "E-", "G-") or rif_clean[0] in "JVGE"):
            continue
        nombre_str = str(nombre).strip()
        if len(nombre_str) < 3:
            continue
        proveedores.append({
            "rif": rif_clean,
            "nombre": nombre_str,
            "direccion": str(ws.cell(row=row, column=4).value or "").strip(),
            "telefono": str(ws.cell(row=row, column=5).value or "").strip(),
            "email": str(ws.cell(row=row, column=6).value or "").strip(),
        })

    # Tambien extraer de TLF DE CONTACTO (puede tener emails)
    if "TLF DE CONTACTO" in wb.sheetnames:
        ws2 = wb["TLF DE CONTACTO"]
        print(f"Hoja: TLF DE CONTACTO ({ws2.max_row} filas)")
        for row in range(2, ws2.max_row + 1):
            rif = ws2.cell(row=row, column=1).value
            email = ws2.cell(row=row, column=2).value
            tel = ws2.cell(row=row, column=3).value
            if rif and email:
                rif_clean = str(rif).strip().upper()
                for p in proveedores:
                    if p["rif"] in rif_clean or rif_clean in p["rif"]:
                        if not p["email"] and "@" in str(email):
                            p["email"] = str(email).strip()
                        if not p["telefono"] and tel:
                            p["telefono"] = str(tel).strip()

    # Deduplicar
    by_rif = {}
    for p in proveedores:
        rif = p["rif"]
        if rif not in by_rif:
            by_rif[rif] = p
        else:
            # Merge: si el nuevo tiene email/telefono, agregarlo
            for k in ["email", "telefono"]:
                if not by_rif[rif].get(k) and p.get(k):
                    by_rif[rif][k] = p[k]

    unicos = sorted(by_rif.values(), key=lambda x: x["nombre"])
    print(f"\nTotal proveedores unicos: {len(unicos)}")
    for p in unicos[:20]:
        e = p.get("email", "") or "sin email"
        print(f"  {p['rif']:20s} | {p['nombre'][:50]:50s} | {e[:30]}")
    if len(unicos) > 20:
        print(f"  ... y {len(unicos)-20} mas")

    # Guardar
    OUT.write_text(json.dumps({
        "fecha_extraccion": "2026-10-13",
        "fuente": FILE.name,
        "total": len(unicos),
        "proveedores": unicos
    }, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nGuardado: {OUT}")


if __name__ == "__main__":
    main()
